"""Builds the public PICO legal site (legal-site/*.html) and optional Word copies.

Usage: python3 docs/legal/build_legal.py [--docx OUT_DIR]
"""
import html, re, sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SITE = ROOT / "legal-site"
EMAIL = "redreishilife+pico@gmail.com"
COMPANY = "RED REISHI LLC"
ADDRESS = "7201 10th Ave N, Saint Petersburg, FL 33710, USA"
UPDATED = "October 25, 2026"

# Blocks: ("h2", text) | ("p", text) | ("ul", [items]) | ("note", text). Inline: **bold**, [text](url)
PRIVACY = ("Privacy Policy", [
    ("p", f"This Privacy Policy explains how {COMPANY} (\"we\", \"us\") handles information when you use the PICO mobile app and related services (\"PICO\"). PICO sells travel eSIM data plans. We built PICO to collect as little as possible: **there are no PICO accounts, no advertising, no tracking and no location collection.**"),
    ("h2", "1. Information we collect"),
    ("ul", [
        "**Order information** — the destination and data plan you choose, price, currency, order number, order status and timestamps.",
        "**eSIM details** — the eSIM activation (installation) code and identifiers issued for your order, so you can install and manage the plan.",
        "**Payment references** — an identifier for your payment from our payment processor. **We never receive or store your full card number.** Card details are entered into and processed by Stripe.",
        "**Referral information** — if you use or share a referral code, the code and the credit it generated.",
        "**Support messages** — if you email us, your email address and anything you include in your message.",
        "**Technical data** — like every internet service, our servers and hosting provider receive your IP address and basic request information (time, app version, error details) in short-lived security logs.",
    ]),
    ("p", "PICO does **not** collect your name, phone number, contacts, photos, precise or approximate location, advertising ID, or browsing activity. PICO does not use analytics or advertising SDKs."),
    ("h2", "2. Information stored on your device"),
    ("p", "Your order numbers and plan details are saved on your phone so the **My Plans** tab can show them without an account. Deleting the app removes this information from your device."),
    ("h2", "3. How we use information"),
    ("ul", [
        "To process your purchase and deliver your eSIM.",
        "To show your plans, data allowance and expiry dates.",
        "To answer support requests, process refunds and apply referral credits.",
        "To prevent fraud and abuse, keep the service secure, and meet legal, tax and accounting obligations.",
    ]),
    ("p", "We do **not** sell or rent your personal information, and we do not share it for cross-context behavioral advertising."),
    ("h2", "4. Who we share it with"),
    ("ul", [
        "**Stripe** — payment processing ([stripe.com/privacy](https://stripe.com/privacy)).",
        "**eSIM network partners** — the wholesale provider that issues your eSIM receives the order reference and plan needed to activate it. Your mobile data traffic is carried by local mobile networks under their own policies.",
        "**Hosting provider** — our servers run on Render ([render.com/privacy](https://render.com/privacy)) in the United States.",
        "**Authorities** — when required by law, or to protect the rights and safety of our users or the public.",
    ]),
    ("h2", "5. How long we keep it"),
    ("p", "We keep order and payment records for as long as needed to provide your plan and for up to 7 years where tax and accounting rules require it. Security logs are kept for up to 30 days. Support emails are kept for up to 2 years after the conversation ends."),
    ("h2", "6. Your choices and rights"),
    ("p", f"You can ask us to access, correct, export or delete information about your orders by emailing [{EMAIL}](mailto:{EMAIL}) with your order number. Depending on where you live (for example the EU/UK under GDPR, or California under the CCPA/CPRA), you may have additional rights, including the right to object to processing and to complain to your local data protection authority. We will not discriminate against you for exercising your rights. We respond within 30 days."),
    ("p", "Legal bases (EU/UK): performance of a contract (delivering your plan), legal obligation (tax records), and legitimate interests (security and fraud prevention)."),
    ("h2", "7. Security"),
    ("p", "Data is encrypted in transit (HTTPS). Payment data is handled by Stripe, a PCI DSS Level 1 provider. Access to our systems is limited and protected. No system is perfectly secure, so please keep your device protected."),
    ("h2", "8. International transfers"),
    ("p", "PICO is operated from the United States. If you use PICO from another country, your information is transferred to and processed in the United States with appropriate safeguards."),
    ("h2", "9. Children"),
    ("p", "PICO is not directed to children under 13, and purchases require the account holder of the payment method. We do not knowingly collect information from children under 13. If you believe a child has given us information, contact us and we will delete it."),
    ("h2", "10. Changes"),
    ("p", "We will post any changes on this page and update the date above. Material changes will also be announced in the app."),
    ("h2", "11. Contact"),
    ("p", f"{COMPANY}<br>{ADDRESS}<br>Email: [{EMAIL}](mailto:{EMAIL})"),
])

TERMS = ("Terms of Service", [
    ("p", f"These Terms of Service (\"Terms\") are an agreement between you and {COMPANY} (\"we\", \"us\") for use of the PICO app and the eSIM data plans sold through it. By using PICO or buying a plan, you agree to these Terms and to our [Privacy Policy](privacy.html)."),
    ("h2", "1. The service"),
    ("p", "PICO lets you buy prepaid mobile **data** plans delivered as an eSIM for use while travelling. Plans are supplied by third-party wholesale eSIM providers and carried on local mobile networks. Unless a plan clearly says otherwise, plans are **data-only** and do not include a phone number, voice calls or SMS."),
    ("note", "Test catalog: while the app shows a \"Test catalog\" banner, plans are simulated, no eSIM is issued and no real payment is taken."),
    ("h2", "2. Eligibility and your device"),
    ("ul", [
        "You must be at least 18, or have permission from a parent or guardian, and be authorized to use the payment method.",
        "Your phone must support eSIM and be **carrier-unlocked**. You confirm compatibility before checkout. We are not responsible for plans bought for incompatible or locked devices.",
        "An internet connection (for example Wi-Fi) is needed to install the eSIM.",
    ]),
    ("h2", "3. Plans, validity and coverage"),
    ("ul", [
        "Each plan lists its countries, data allowance and validity period. Validity normally starts when the eSIM first connects to a supported network in a covered country.",
        "Coverage, speed (including 4G/5G availability) and network choice depend on local operators and can vary. We do not guarantee uninterrupted or error-free service.",
        "Unused data expires at the end of the validity period and is not refundable or transferable unless the plan says otherwise.",
        "Some eSIMs can only be installed once. Do not delete an installed eSIM from your phone unless you no longer need it.",
    ]),
    ("h2", "4. No emergency calls"),
    ("p", "**PICO data plans cannot be used to call emergency services (such as 911, 112 or 999).** Do not rely on PICO for emergency communications. Keep another way to reach emergency services when travelling."),
    ("h2", "5. Prices and payment"),
    ("p", "Prices are shown in the app before you pay and include any fees we charge. Taxes may apply depending on your location. Payments are processed by Stripe. Your purchase is complete when payment is confirmed and the order appears in **My Plans**."),
    ("h2", "6. Refunds"),
    ("ul", [
        "**Not yet installed:** you can request a full refund within 30 days of purchase if the eSIM has not been installed or activated.",
        "**Installed but not working:** if the eSIM was installed on a compatible, unlocked device but could not connect in a covered country, and our support steps did not fix it, we will refund the unused plan.",
        "Refunds are not available for data already used, for incompatible or locked devices, or for plans that have expired.",
        f"To request a refund, email [{EMAIL}](mailto:{EMAIL}) with your order number. Approved refunds go back to the original payment method. This does not affect any rights you have under consumer protection law.",
    ]),
    ("h2", "7. Referral credits"),
    ("p", "If PICO offers referral credits, the amount and maximum are shown in the app. Credits have no cash value, cannot be sold or transferred, and may be removed if earned through fraud, self-referral or abuse. We may change or end the referral program at any time; earned credits remain usable for 12 months."),
    ("h2", "8. Acceptable use"),
    ("p", "You agree not to use PICO or any plan for fraud, phishing, spam, harassment, malware, illegal content, money laundering, sanctions evasion, network attacks, reselling plans without our permission, or anything unlawful. We and our network partners may suspend or end service that breaks these rules or local law, without refund where permitted."),
    ("h2", "9. Third-party services"),
    ("p", "Network service is provided by third-party operators and eSIM providers, and payments by Stripe. Their own terms may also apply. We are not responsible for outages or actions of these third parties, but we will help you resolve problems."),
    ("h2", "10. Disclaimers and liability"),
    ("p", "PICO is provided \"as is\" and \"as available\" to the fullest extent permitted by law. To the extent permitted by law, we are not liable for indirect, incidental, special or consequential damages, or lost profits or data, and our total liability for any claim is limited to the amount you paid for the plan the claim relates to. Some places do not allow these limits, so they may not apply to you."),
    ("h2", "11. Changes and termination"),
    ("p", "We may update these Terms; the new version applies from the date shown above. Changes do not affect plans you already bought. You can stop using PICO at any time by deleting the app."),
    ("h2", "12. Apple App Store"),
    ("p", "If you downloaded PICO from the Apple App Store, these Terms are between you and us, not Apple. Apple has no obligation to provide maintenance or support, is not responsible for claims relating to the app, and is a third-party beneficiary of these Terms with the right to enforce them against you."),
    ("h2", "13. Governing law"),
    ("p", "These Terms are governed by the laws of the State of Florida, USA, without regard to conflict-of-laws rules. Disputes will be handled in the state or federal courts located in Pinellas County, Florida, unless your local consumer law gives you the right to bring a claim where you live."),
    ("h2", "14. Contact"),
    ("p", f"{COMPANY}<br>{ADDRESS}<br>Email: [{EMAIL}](mailto:{EMAIL})"),
])

SUPPORT = ("PICO Support", [
    ("p", f"Need help with your eSIM? Email **[{EMAIL}](mailto:{EMAIL})** — we reply within 1 business day. Please include your order number (shown in **My Plans**) and your phone model."),
    ("h2", "Before you buy"),
    ("ul", [
        "**Is my phone compatible?** Your phone must support eSIM and be carrier-unlocked. iPhone: Settings → General → About — look for \"Digital SIM\" or \"Available SIM\". Android: Settings → Connections/Network → SIM manager — look for \"Add eSIM\".",
        "**Does it include a phone number?** No. PICO plans are data-only unless the plan says otherwise. You can still use apps like WhatsApp or FaceTime over data.",
        "**When does my plan start?** When the eSIM first connects to a network in a covered country — not when you buy it.",
    ]),
    ("h2", "Installing your eSIM"),
    ("ul", [
        "Install while you have Wi-Fi, ideally before your trip.",
        "Open **My Plans**, choose your plan and follow the install steps, or scan the QR code from another screen.",
        "Do not delete the eSIM after installing — some eSIMs cannot be installed twice.",
    ]),
    ("h2", "I can't connect"),
    ("ul", [
        "Make sure the PICO eSIM line is turned **on** and selected for **Mobile/Cellular data**.",
        "Turn **Data Roaming on** for the PICO eSIM line (this does not create roaming charges from your home carrier when your home line's data is off).",
        "Toggle Airplane mode on and off, or restart your phone.",
        "Select a network manually: Settings → Mobile network → Network operators → choose a listed partner network.",
        "Check you are in a country covered by your plan and that the plan has not expired.",
        "Still stuck? Open **Help → I can't connect** in the app, or email us.",
    ]),
    ("h2", "Refunds"),
    ("p", "Unused, uninstalled plans can be refunded within 30 days. If your eSIM is installed but won't connect and our steps don't fix it, we'll refund the unused plan. See the [Terms of Service](terms.html#6-refunds) for details."),
    ("h2", "Your data"),
    ("p", f"PICO has no accounts and no tracking. To access or delete your order information, email [{EMAIL}](mailto:{EMAIL}) with your order number. Read the [Privacy Policy](privacy.html)."),
    ("h2", "Emergency calls"),
    ("p", "PICO data plans **cannot call emergency services**. Keep another way to reach 911/112/999 while travelling."),
    ("h2", "Company"),
    ("p", f"PICO is operated by {COMPANY}, {ADDRESS}."),
])

PAGES = {"privacy.html": PRIVACY, "terms.html": TERMS, "support.html": SUPPORT}

CSS = """*{box-sizing:border-box}body{margin:0;font:17px/1.6 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Helvetica,Arial,sans-serif;color:#0B1B3F;background:#fff}
header{background:#00175B;color:#fff;padding:18px 20px}header .in{max-width:820px;margin:auto;display:flex;align-items:center;gap:14px}
header img{width:48px;height:48px;border-radius:12px}header a{color:#fff;text-decoration:none;font-weight:800;font-size:22px;letter-spacing:.5px}
nav{max-width:820px;margin:auto;padding:10px 20px;display:flex;gap:18px;flex-wrap:wrap;border-bottom:1px solid #E3E8F2}nav a{color:#1D5FD6;text-decoration:none;font-weight:600}
main{max-width:820px;margin:auto;padding:8px 20px 48px}h1{font-size:30px;margin:24px 0 4px}h2{font-size:20px;margin:28px 0 6px}
.upd{color:#5A6B8C;font-size:14px;margin:0 0 18px}a{color:#1D5FD6}li{margin:6px 0}
.note{background:#FFF7E6;border:1px solid #F0C36D;border-radius:10px;padding:10px 14px;color:#7A4E00}
footer{background:#F4F6FB;color:#5A6B8C;font-size:14px;text-align:center;padding:20px}"""


def slug(t):
    return re.sub(r"[^a-z0-9]+", "-", t.lower()).strip("-")


def inline_html(t):
    t = html.escape(t, quote=False).replace("&lt;br&gt;", "<br>")
    t = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", t)
    return re.sub(r"\[(.+?)\]\((.+?)\)", r'<a href="\2">\1</a>', t)


def page_html(fname, title, blocks):
    body = []
    for kind, val in blocks:
        if kind == "h2":
            body.append(f'<h2 id="{slug(val)}">{inline_html(val)}</h2>')
        elif kind == "ul":
            body.append("<ul>" + "".join(f"<li>{inline_html(i)}</li>" for i in val) + "</ul>")
        elif kind == "note":
            body.append(f'<p class="note">{inline_html(val)}</p>')
        else:
            body.append(f"<p>{inline_html(val)}</p>")
    navs = " ".join(f'<a href="{f}">{PAGES[f][0].replace("PICO ", "")}</a>' for f in PAGES)
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>{title} — PICO</title>
<link rel="icon" href="icon.png"><style>{CSS}</style></head><body>
<header><div class="in"><img src="icon.png" alt="PICO"><a href="index.html">PICO</a></div></header>
<nav>{navs}<a href="mailto:{EMAIL}">Contact</a></nav>
<main><h1>{title}</h1><p class="upd">Last updated: {UPDATED}</p>
{chr(10).join(body)}</main>
<footer>&copy; 2026 {COMPANY} · {ADDRESS} · <a href="mailto:{EMAIL}">{EMAIL}</a></footer></body></html>
"""


def build_site():
    SITE.mkdir(exist_ok=True)
    for fname, (title, blocks) in PAGES.items():
        (SITE / fname).write_text(page_html(fname, title, blocks))
    (SITE / "index.html").write_text(page_html("index.html", "PICO — travel eSIM data", [
        ("p", "PICO gives you mobile data abroad in minutes: pick your destination, buy a data plan, install the eSIM and connect when you land."),
        ("ul", ["[Support & FAQ](support.html)", "[Privacy Policy](privacy.html)", "[Terms of Service](terms.html)",
                f"Contact: [{EMAIL}](mailto:{EMAIL})"]),
    ]))
    from PIL import Image
    Image.open(ROOT / "app/assets/APP_PICO.png").resize((192, 192), Image.LANCZOS).save(SITE / "icon.png")


def build_docx(out):
    from docx import Document
    from docx.shared import Pt
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)

    def add_runs(par, text):
        text = re.sub(r"\[(.+?)\]\((.+?)\)", lambda m: m.group(1) if m.group(1) == m.group(2) or "@" in m.group(1) or "." in m.group(1) else f"{m.group(1)} ({m.group(2)})", text)
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            for j, line in enumerate(part.split("<br>")):
                if j:
                    par.add_run().add_break()
                par.add_run(line).bold = bool(i % 2)

    for fname, (title, blocks) in PAGES.items():
        d = Document()
        d.styles["Normal"].font.size = Pt(11)
        d.add_heading(f"{title} — PICO", 0)
        d.add_paragraph(f"Last updated: {UPDATED}")
        for kind, val in blocks:
            if kind == "h2":
                d.add_heading(val, 1)
            elif kind == "ul":
                for item in val:
                    add_runs(d.add_paragraph(style="List Bullet"), item)
            else:
                add_runs(d.add_paragraph(), val)
        d.save(out / f"PICO_{title.replace('PICO ', '').replace(' ', '_')}.docx")


if __name__ == "__main__":
    build_site()
    if "--docx" in sys.argv:
        build_docx(sys.argv[sys.argv.index("--docx") + 1])
